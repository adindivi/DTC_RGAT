# -*- coding: utf-8 -*-
"""
DTC 지식 그래프 및 엑셀 데이터 무결성 검증 모듈
- 단일 책임 원칙(SRP) 준수 및 함수형 모듈화
- ISO 14229-1 FTB 서브타입 빈칸 및 '미분류' 전수 검증
- Windows 콘솔 인코딩(CP949) 호환 안전 출력
- pytest 테스트 및 독립 CLI 스크립트 양방향 지원
"""
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

import config

BASE_DIR = Path(__file__).parent.resolve()
DEFAULT_JSON_PATH = BASE_DIR / "dtc_knowledge_graph.json"
DEFAULT_EXCEL_PATH = BASE_DIR / "dtc_knowledge_graph.xlsx"

SAMPLE_TEST_CODES = [
    "C133A00", "B151114", "B132656", "P0D0972",
    "P0D2498", "C120102", "C139846", "B100252",
]


@dataclass
class ValidationSummary:
    """데이터 무결성 검증 결과 요약 데이터 클래스"""
    total_nodes: int
    empty_subtypes: int
    unclassified_subtypes: int
    is_valid: bool


def validate_json_subtypes(json_path: Path = DEFAULT_JSON_PATH) -> tuple[ValidationSummary, list[dict[str, Any]]]:
    """
    JSON 파일 내 DTC 노드의 서브타입 명칭 무결성 검증 (SRP: JSON 데이터 검증)
    """
    if not json_path.exists():
        raise FileNotFoundError(f"JSON graph file not found: {json_path}")

    with open(json_path, "r", encoding="utf-8") as f:
        graph_data = json.load(f)

    dtc_nodes = [
        node for node in graph_data.get("nodes", [])
        if node.get("node_type") == "DTC"
    ]
    empty_nodes = [
        node for node in dtc_nodes
        if not (node.get("subtype_name") or "").strip()
    ]
    unclassified_nodes = [
        node for node in dtc_nodes
        if (node.get("subtype_name") or "").strip() == "미분류"
    ]

    summary = ValidationSummary(
        total_nodes=len(dtc_nodes),
        empty_subtypes=len(empty_nodes),
        unclassified_subtypes=len(unclassified_nodes),
        is_valid=(len(empty_nodes) == 0 and len(unclassified_nodes) == 0)
    )
    return summary, dtc_nodes


def validate_excel_subtypes(excel_path: Path = DEFAULT_EXCEL_PATH) -> ValidationSummary:
    """
    엑셀 마스터 파일의 7번째 열(서브타입 명칭) 무결성 검증 (SRP: 엑셀 데이터 검증)
    """
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel master file not found: {excel_path}")

    workbook = openpyxl.load_workbook(str(excel_path), data_only=True, read_only=True)
    target_sheet = "2. DTC 종합 마스터"
    if target_sheet not in workbook.sheetnames:
        raise KeyError(f"Sheet '{target_sheet}' not found in {excel_path.name}")

    worksheet = workbook[target_sheet]

    total_rows = 0
    empty_rows = 0
    unclassified_rows = 0

    for row in worksheet.iter_rows(min_row=2, min_col=1, max_col=12, values_only=True):
        total_rows += 1
        subtype_name = str(row[6] or "").strip()
        if not subtype_name:
            empty_rows += 1
        elif subtype_name == "미분류":
            unclassified_rows += 1

    workbook.close()

    return ValidationSummary(
        total_nodes=total_rows,
        empty_subtypes=empty_rows,
        unclassified_subtypes=unclassified_rows,
        is_valid=(empty_rows == 0 and unclassified_rows == 0)
    )


def verify_sample_codes(dtc_nodes: list[dict[str, Any]], sample_codes: list[str]) -> list[dict[str, str]]:
    """
    대표 샘플 DTC 코드의 FTB 코드 및 서브타입 명칭 매핑 확인
    """
    node_by_code = {node.get("code"): node for node in dtc_nodes if node.get("code")}
    results = []
    for code in sample_codes:
        node = node_by_code.get(code)
        if node:
            results.append({
                "code": code,
                "subtype_code": node.get("subtype_code", ""),
                "subtype_name": node.get("subtype_name", ""),
            })
    return results


def run_full_validation(
    json_path: Path = DEFAULT_JSON_PATH,
    excel_path: Path = DEFAULT_EXCEL_PATH
) -> bool:
    """
    전체 무결성 검증 오케스트레이터 및 콘솔 리포트 출력
    """
    # Windows CP949 인코딩 안전 출력 설정
    if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except AttributeError:
            pass

    print("=" * 60)
    print(f" 1. JSON 무결성 검증 ({json_path.name})")
    print("=" * 60)
    json_summary, dtc_nodes = validate_json_subtypes(json_path)
    print(f"  - 총 DTC 노드 수       : {json_summary.total_nodes}개")
    print(f"  - 빈칸('') 서브타입 수 : {json_summary.empty_subtypes}개 (목표: 0)")
    print(f"  - '미분류' 서브타입 수 : {json_summary.unclassified_subtypes}개 (목표: 0)")

    print("\n" + "=" * 60)
    print(f" 2. 엑셀 파일 무결성 검증 ({excel_path.name})")
    print("=" * 60)
    excel_summary = validate_excel_subtypes(excel_path)
    print(f"  - 총 마스터 데이터 행   : {excel_summary.total_nodes}개")
    print(f"  - 엑셀 내 빈칸 서브타입 : {excel_summary.empty_subtypes}개 (목표: 0)")
    print(f"  - 엑셀 내 '미분류'      : {excel_summary.unclassified_subtypes}개 (목표: 0)")

    print("\n" + "=" * 60)
    print(" 3. 신규 매핑 대표 샘플 확인 (상위 8건)")
    print("=" * 60)
    samples = verify_sample_codes(dtc_nodes, SAMPLE_TEST_CODES)
    for sample in samples:
        print(f"  [{sample['code']}] FTB: {sample['subtype_code']:<2} -> 서브타입명: '{sample['subtype_name']}'")

    all_passed = json_summary.is_valid and excel_summary.is_valid

    print("\n" + "=" * 60)
    if all_passed:
        print(" [SUCCESS] JSON 및 Excel 전수 무결성 검증 100% 정상 통과!")
    else:
        print(" [FAILURE] 데이터 무결성 검증 실패 (빈칸 또는 미분류 데이터 존재)")
    print("=" * 60)

    return all_passed


if __name__ == "__main__":
    success = run_full_validation()
    sys.exit(0 if success else 1)
